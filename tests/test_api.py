import datetime
import io
import sys
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'backend'))
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError
import main


class TransportTests(unittest.TestCase):
    def test_required_fields_rejected_by_api_schema(self):
        with self.assertRaises(ValidationError):
            main.JobIn(deposit_date='2026-10-15', delivery_date='2026-10-15')

    def test_export_preserves_text_and_filters(self):
        row = dict(seq=7, job_no='TP-0007', deposit_date='2026-10-15',
                   delivery_date='2026-10-16', to_name='ผู้รับ', to_dept='',
                   to_branch_code='PG', to_branch_name='Paragon', from_name='ผู้ส่ง',
                   from_dept='', from_branch_code='HQ', item_detail='เอกสาร',
                   qty=21, unit='ซอง', remark='=1+1')
        with patch.object(main, 'list_jobs', return_value={'items': [row]}):
            with TestClient(main.app) as client:
                response = client.get('/api/export?date=2026-10-15&q=PG')
                self.assertEqual(response.status_code, 200)
                ws = load_workbook(io.BytesIO(response.content)).active
                self.assertEqual(ws.max_row, 2)
                self.assertEqual(ws['A2'].value, 7)
                self.assertEqual(ws['L2'].value, 21)
                self.assertEqual(ws['N2'].value, '=1+1')
                self.assertEqual(ws['N2'].data_type, 's')
                empty = client.get('/api/export?date=2026-10-15&q=missing')
                self.assertEqual(load_workbook(io.BytesIO(empty.content)).active.max_row, 1)

    def test_job_numbers_and_quantity_bounds(self):
        self.assertEqual(main.job_no(1), 'TP-0001')
        self.assertEqual(main.job_no(10000), 'TP-10000')
        valid = dict(deposit_date=datetime.date.today(), delivery_date=datetime.date.today(),
                     to_branch_code='PG', from_name='ผู้ส่ง', item_detail='เอกสาร')
        for override in [{'qty': 0}, {'qty': 100000}, {'qty': 1.2}, {'unit': 'ลัง'}]:
            with self.assertRaises(ValidationError):
                main.JobIn(**valid, **override)


if __name__ == '__main__':
    unittest.main()
