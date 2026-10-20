import base64
import datetime
import io
import json
import os
import sys
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import send_daily_mail as mail

ROW = dict(id=1, seq=7, job_no='TP-0007', work_date='2026-10-19', deposit_date='2026-10-19',
           delivery_date='2026-10-20', to_name='คุณ<b>เอ</b>', to_dept='คลัง', to_branch_code='PG',
           to_branch_name='Siam Paragon', to_brand='GH', to_shop_id='1', from_name='ผู้ส่ง',
           from_dept='', from_branch_code='', from_branch_name='', from_brand='', from_shop_id='',
           item_detail='เอกสาร', qty=21, unit='ซอง', remark='<script>x</script>',
           created_at='', updated_at='')
ENV = {'MAIL_MODE': 'test', 'MAIL_TEST_TO': 'tester@example.com', 'MAIL_TO': 'a@example.com, b@example.com',
       'MAIL_FROM': 'admin@example.com', 'MAIL_SENDER': 'graph',
       'MS_TENANT_ID': 'tenant', 'MS_CLIENT_ID': 'client', 'MS_CLIENT_SECRET': 'secret'}


class FakeResponse:
    def __init__(self, body=b''):
        self.body = body
    def __enter__(self):
        return self
    def __exit__(self, *args):
        pass
    def read(self):
        return self.body


class MailTests(unittest.TestCase):
    def test_midnight_run_sends_previous_day(self):
        sent = {}
        with patch.dict(os.environ, ENV), \
             patch.object(mail.app, 'now_th', return_value=datetime.datetime(2026, 10, 20, 0, 0, 5)), \
             patch.object(mail.app, 'list_jobs', side_effect=lambda d: sent.setdefault('date', d) and {'items': []}), \
             patch.object(mail, 'already_sent', return_value=False), \
             patch.object(mail, 'log_attempt'), \
             patch.dict(mail.SENDERS, {'graph': lambda *a: sent.setdefault('to', a[1])}):
            self.assertEqual(mail.run([]), 0)
        self.assertEqual(sent['date'], datetime.date(2026, 10, 19))
        self.assertEqual(sent['to'], ['tester@example.com'])        # test mode → ผู้รับทดสอบเท่านั้น

    def test_live_mode_uses_real_recipients_and_no_test_tag(self):
        calls = []
        with patch.dict(os.environ, {**ENV, 'MAIL_MODE': 'live'}), \
             patch.object(mail.app, 'list_jobs', return_value={'items': [ROW]}), \
             patch.object(mail, 'already_sent', return_value=False), \
             patch.object(mail, 'log_attempt'), \
             patch.dict(mail.SENDERS, {'graph': lambda *a: calls.append(a)}):
            self.assertEqual(mail.run(['--date', '2026-10-19']), 0)
        sender, to, subject, body, attachments = calls[0]
        self.assertEqual(to, ['a@example.com', 'b@example.com'])
        self.assertEqual(subject, 'ใบแจ้งงานขนส่ง SST ประจำวันที่ 19/10/2026 (1 รายการ)')
        self.assertNotIn('อีเมลทดสอบ', body)
        self.assertEqual(attachments[0][0], 'ใบแจ้งงานขนส่ง_2026-10-19.xlsx')

    def test_already_sent_day_is_skipped_but_force_and_manual_send(self):
        calls = []
        with patch.dict(os.environ, ENV), \
             patch.object(mail.app, 'list_jobs', return_value={'items': []}), \
             patch.object(mail, 'already_sent', return_value=True), \
             patch.object(mail, 'log_attempt'), \
             patch.dict(mail.SENDERS, {'graph': lambda *a: calls.append(a[1])}):
            mail.run(['--date', '2026-10-19'])
            self.assertEqual(calls, [])                                  # รอบ retry ไม่ส่งซ้ำ
            mail.run(['--date', '2026-10-19', '--force'])
            mail.run(['--date', '2026-10-19', '--to', 'x@example.com'])
        self.assertEqual(calls, [['tester@example.com'], ['x@example.com']])

    def test_failure_is_logged_and_returns_error(self):
        def boom(*args):
            raise RuntimeError('ยังไม่ได้ตั้งค่า')
        with patch.dict(os.environ, ENV), \
             patch.object(mail.app, 'list_jobs', return_value={'items': []}), \
             patch.object(mail, 'already_sent', return_value=False), \
             patch.object(mail, 'log_attempt') as log, \
             patch.dict(mail.SENDERS, {'graph': boom}):
            self.assertEqual(mail.run(['--date', '2026-10-19']), 1)
        self.assertEqual(log.call_args.args[4], 'failed')

    def run_failing(self, args, list_jobs=None):
        def boom(*a):
            raise RuntimeError('HTTP 401: invalid_client <b>')
        alerts = []
        with patch.dict(os.environ, {**ENV, 'MAIL_ALERT_TO': 'it1@example.com, it2@example.com'}), \
             patch.object(mail.app, 'list_jobs', **(list_jobs or {'return_value': {'items': []}})), \
             patch.object(mail, 'already_sent', return_value=False), \
             patch.object(mail, 'log_attempt'), \
             patch.object(mail, 'send_alert', side_effect=lambda *a, **k: alerts.append(a) or True), \
             patch.dict(mail.SENDERS, {'graph': boom}):
            code = mail.run(args)
        return code, alerts

    def test_only_the_final_round_alerts_it_on_failure(self):
        self.assertEqual(self.run_failing(['--date', '2026-10-19']), (1, []))                 # 00:00 / 00:10
        code, alerts = self.run_failing(['--date', '2026-10-19', '--final'])                   # 00:30
        self.assertEqual(code, 1)
        self.assertEqual(alerts[0][0], datetime.date(2026, 10, 19))
        self.assertIn('invalid_client', alerts[0][1])

    def test_database_down_is_caught_and_alerted(self):
        code, alerts = self.run_failing(['--date', '2026-10-19', '--final'],
                                        list_jobs={'side_effect': OSError('MySQL down')})
        self.assertEqual(code, 1)
        self.assertIn('MySQL down', alerts[0][1])

    def test_success_or_skip_on_final_round_does_not_alert(self):
        alerts = []
        with patch.dict(os.environ, ENV), \
             patch.object(mail.app, 'list_jobs', return_value={'items': []}), \
             patch.object(mail, 'log_attempt'), \
             patch.object(mail, 'send_alert', side_effect=lambda *a, **k: alerts.append(a)), \
             patch.dict(mail.SENDERS, {'graph': lambda *a: None}):
            with patch.object(mail, 'already_sent', return_value=False):
                self.assertEqual(mail.run(['--date', '2026-10-19', '--final']), 0)
            with patch.object(mail, 'already_sent', return_value=True):
                self.assertEqual(mail.run(['--date', '2026-10-19', '--final']), 0)
        self.assertEqual(alerts, [])

    def test_alert_goes_to_it_recipients_with_escaped_error(self):
        sent = []
        with patch.dict(os.environ, {**ENV, 'MAIL_ALERT_TO': 'it1@example.com; it2@example.com'}), \
             patch.object(mail, 'log_attempt') as log, \
             patch.dict(mail.SENDERS, {'graph': lambda *a: sent.append(a)}):
            self.assertTrue(mail.send_alert(datetime.date(2026, 10, 19), 'boom <script>', '2026-10-20 00:30:01'))
        sender, to, subject, body, attachments = sent[0]
        self.assertEqual(to, ['it1@example.com', 'it2@example.com'])
        self.assertIn('[แจ้งเตือน]', subject)
        self.assertIn('19/10/2026', subject)
        self.assertIn('boom &lt;script&gt;', body)
        self.assertIn('--date 2026-10-19', body)                              # คำสั่งส่งซ้ำ
        self.assertEqual(attachments, [])
        self.assertEqual(log.call_args.args[1], 'alert')

    def test_alert_channel_failure_does_not_crash(self):
        def boom(*a):
            raise RuntimeError('graph down')
        with patch.dict(os.environ, {**ENV, 'MAIL_ALERT_TO': 'it@example.com'}), \
             patch.object(mail, 'log_attempt'), patch.dict(mail.SENDERS, {'graph': boom}):
            self.assertFalse(mail.send_alert(datetime.date(2026, 10, 19), 'x', 'now'))
        with patch.dict(os.environ, {**ENV, 'MAIL_ALERT_TO': ''}):
            self.assertFalse(mail.send_alert(datetime.date(2026, 10, 19), 'x', 'now'))

    def test_html_escapes_user_text_and_handles_empty_day(self):
        body = mail.build_html(datetime.date(2026, 10, 19), [ROW], 'test')
        self.assertIn('&lt;script&gt;x&lt;/script&gt;', body)
        self.assertNotIn('<script>x', body)
        self.assertIn('คุณ&lt;b&gt;เอ&lt;/b&gt;', body)
        self.assertIn('วันจันทร์ที่ 19 ตุลาคม 2026', body)
        self.assertIn('ซอง 21', body)
        self.assertIn('อีเมลทดสอบ', body)
        self.assertIn('กรุณาอย่าตอบกลับ', body)                        # แจ้งเมลส่งอย่างเดียว
        empty = mail.build_html(datetime.date(2026, 10, 17), [], 'live')
        self.assertIn('ไม่มีรายการใบแจ้งงานขนส่ง', empty)
        self.assertNotIn('<th', empty)

    def test_graph_request_sends_html_and_excel_attachment(self):
        requests = []
        def fake_urlopen(req, timeout=None):
            requests.append(req)
            if 'oauth2' in req.full_url:
                return FakeResponse(json.dumps({'access_token': 'tok'}).encode())
            return FakeResponse(b'')                                      # sendMail ตอบ 202 ไม่มี body
        xlsx = mail.app.build_xlsx([ROW])
        with patch.dict(os.environ, ENV), patch.object(mail.urllib.request, 'urlopen', fake_urlopen):
            mail.send_graph('admin@example.com', ['a@example.com'], 'หัวเรื่อง', '<p>x</p>',
                            [('f.xlsx', mail.XLSX_TYPE, xlsx)])
        token_req, send_req = requests
        self.assertIn('/tenant/oauth2/v2.0/token', token_req.full_url)
        self.assertIn(b'grant_type=client_credentials', token_req.data)
        self.assertEqual(send_req.full_url, 'https://graph.microsoft.com/v1.0/users/admin%40example.com/sendMail')
        self.assertEqual(send_req.get_header('Authorization'), 'Bearer tok')
        msg = json.loads(send_req.data)['message']
        self.assertEqual(msg['toRecipients'], [{'emailAddress': {'address': 'a@example.com'}}])
        self.assertEqual(msg['body']['contentType'], 'HTML')
        ws = load_workbook(io.BytesIO(base64.b64decode(msg['attachments'][0]['contentBytes']))).active
        self.assertEqual(ws['B2'].value, 'TP-0007')

    def test_graph_without_credentials_explains_what_is_missing(self):
        with patch.dict(os.environ, {**ENV, 'MS_CLIENT_SECRET': ''}):
            with self.assertRaisesRegex(RuntimeError, 'MS_CLIENT_SECRET'):
                mail.send_graph('admin@example.com', ['a@example.com'], 's', 'b', [])


if __name__ == '__main__':
    unittest.main()
