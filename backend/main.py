#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# =============================================================================
#  main.py — ใบแจ้งงานขนส่ง (Transport Notice) API — FastAPI + MySQL (transport_notice)
#
#  รันที่ :8606 ภายในสำนักงาน
#  หน้าเว็บจริงอยู่ที่ index.html ในรากโปรเจกต์ → เปิด http://10.8.1.88:8606 ได้เลย
#  GitHub Pages เรียก API ผ่าน Cloudflare Tunnel https://transport-notice.mobile1234.site
#  ตอนนี้เปิดใช้ได้เลยไม่ต้องใช้รหัส; ถ้าใส่ ACCESS_KEY ใน .env คำขอที่มาทาง Tunnel จะต้องมีรหัส
#
#  รัน:  cd backend && ../venv/bin/uvicorn main:app --host 0.0.0.0 --port 8606
#        (หรือ systemd — ดู transport-notice-api.service)
#
#  เลขงาน: ลำดับ/Job No. เริ่ม 1 ใหม่ทุกเที่ยงคืน (เวลาไทย) → TP-0001, TP-0002, ...
# =============================================================================
import datetime
import hmac
import logging
import os
import re
from io import BytesIO
from urllib.parse import quote
from typing import Optional
from zoneinfo import ZoneInfo

from fastapi import FastAPI, HTTPException, Query, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse, RedirectResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, ConfigDict, Field, field_validator

from db import ROOT, get_conn

log = logging.getLogger("transport_notice")

TZ = ZoneInfo("Asia/Bangkok")
JOB_PREFIX = "TP-"
JOB_DIGITS = 4
DEFAULT_UNITS = ["กล่อง", "ซอง", "แพค", "ขวด", "ชิ้น"]   # ใช้เมื่อยังไม่มีตารางหน่วย (หน่วยจริงแก้ในหน้าตั้งค่า)
MANUAL_BRAND = "GHC"                                        # สาขาที่เพิ่มจากหน้าตั้งค่า
QTY_MAX = 99_999
DOCS_DIR = os.path.join(ROOT, "docs")
INDEX_FILE = os.path.join(ROOT, "index.html")

JOB_COLUMNS = [
    "id", "work_date", "seq", "job_no", "deposit_date", "delivery_date",
    "to_name", "to_dept", "to_branch_code", "to_branch_name", "to_brand", "to_shop_id",
    "from_name", "from_dept", "from_branch_code", "from_branch_name", "from_brand", "from_shop_id",
    "item_detail", "qty", "unit", "remark", "created_at", "updated_at",
]
JOB_SELECT = "SELECT " + ", ".join(JOB_COLUMNS) + " FROM transport_notice_jobs"

app = FastAPI(title="Transport Notice API", version="1.1", docs_url="/api/docs",
              redoc_url=None, openapi_url="/api/openapi.json")


def _norm_key(value):
    # ผู้ใช้พิมพ์ตัวเล็ก/เว้นวรรค/ขีดได้ เช่น "k7qm 4xpa r2ht" = "K7QM-4XPA-R2HT"
    return re.sub(r"[\s-]", "", value or "").upper()


ACCESS_KEY = _norm_key(os.environ.get("ACCESS_KEY"))


# ACCESS_KEY ว่าง = ไม่ต้องใช้รหัส (ค่าปัจจุบัน) / ใส่รหัส = คำขอผ่าน Cloudflare Tunnel
# (มี header CF-Connecting-IP = มาจากอินเทอร์เน็ต) ต้องส่ง X-Access-Key ให้ถูก ในออฟฟิศไม่ต้องใช้รหัส
# ประกาศก่อน CORSMiddleware เพื่อให้ CORS ห่อด้านนอก: preflight ผ่าน และ 401 มี CORS header
@app.middleware("http")
async def require_access_key(request: Request, call_next):
    if (ACCESS_KEY and request.url.path.startswith("/api/") and request.method != "OPTIONS"
            and "cf-connecting-ip" in request.headers):
        given = _norm_key(request.headers.get("x-access-key"))
        if not hmac.compare_digest(given.encode(), ACCESS_KEY.encode()):
            return JSONResponse(status_code=401, content={
                "ok": False, "need_key": True,
                "error": "รหัสเข้าใช้งานไม่ถูกต้อง" if given else "กรุณาใส่รหัสเข้าใช้งาน"})
    return await call_next(request)


# รองรับเปิด index.html ในเครื่องร่วมกับ API ภายในสำนักงาน (ไม่มี login/cookie)
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["GET", "POST", "PUT", "DELETE"],
    allow_headers=["*"],
)


# error ที่หลุดออกจาก endpoint จะไม่มี CORS header → browser โชว์เป็น "CORS blocked" หลอกๆ
# ตัวนี้ดักไว้แล้วใส่ header ให้เสมอ (รายละเอียด error จริงดูใน journalctl)
@app.exception_handler(Exception)
async def catch_all(request: Request, exc: Exception):
    log.exception("unhandled error on %s %s", request.method, request.url.path)
    return JSONResponse(
        status_code=500,
        content={"ok": False, "error": f"ระบบขัดข้อง ({type(exc).__name__}) ลองใหม่อีกครั้ง"},
        headers={"Access-Control-Allow-Origin": "*"},
    )


def now_th():
    return datetime.datetime.now(TZ).replace(tzinfo=None)


def job_no(seq):
    return f"{JOB_PREFIX}{seq:0{JOB_DIGITS}d}"


def _out(row):
    # date/datetime → string ให้ JS ใช้ตรงๆ
    for k, v in row.items():
        if isinstance(v, datetime.datetime):
            row[k] = v.strftime("%Y-%m-%d %H:%M:%S")
        elif isinstance(v, datetime.date):
            row[k] = v.isoformat()
    return row


def _get_job(cur, job_id):
    cur.execute(JOB_SELECT + " WHERE id=%s AND deleted_at IS NULL", (job_id,))
    row = cur.fetchone()
    if not row:
        raise HTTPException(404, "ไม่พบรายการนี้ (อาจถูกลบไปแล้ว)")
    return _out(row)


# ---------- Schemas ----------
class JobIn(BaseModel):
    model_config = ConfigDict(str_strip_whitespace=True, validate_default=True)

    deposit_date: datetime.date
    delivery_date: datetime.date
    to_name: str = Field("", max_length=150)
    to_dept: str = Field("", max_length=150)
    to_branch_code: str = Field("", max_length=30)
    to_branch_name: str = Field("", max_length=255)
    to_brand: str = Field("", max_length=10)
    to_shop_id: str = Field("", max_length=20)
    from_name: str = Field("", max_length=150)
    from_dept: str = Field("", max_length=150)
    from_branch_code: str = Field("", max_length=30)
    from_branch_name: str = Field("", max_length=255)
    from_brand: str = Field("", max_length=10)
    from_shop_id: str = Field("", max_length=20)
    item_detail: str = Field("", max_length=500)
    qty: int = Field(1, ge=1, le=QTY_MAX)
    unit: str = Field("กล่อง", min_length=1, max_length=10)   # ตรวจกับหน่วยที่เปิดใช้ตอนบันทึก (create_job)
    remark: str = Field("", max_length=500)

    @field_validator("to_branch_code", "from_name", "item_detail")
    @classmethod
    def _required(cls, v):
        if not v:
            raise ValueError("กรุณากรอกสาขาปลายทาง ผู้ส่ง และรายละเอียด สิ่งที่ส่ง")
        return v

WRITE_FIELDS = list(JobIn.model_fields)


# ---------- Endpoints ----------
@app.get("/api/health")
def health():
    cn = get_conn()
    try:
        with cn.cursor() as cur:
            cur.execute("SELECT COUNT(*) AS n FROM transport_notice_branches WHERE is_active=1")
            n = cur.fetchone()["n"]
    finally:
        cn.close()
    return {"ok": True, "today": now_th().date().isoformat(), "branches": n}


@app.get("/api/meta")
def meta():
    # หน้าเว็บเรียกตอนเปิด — ถ้า MySQL ล่ม (เช่น OOM 16 ต.ค.) ยังได้ "วันนี้" + หน่วยเริ่มต้น หน้าเว็บไม่พังทั้งหน้า
    try:
        cn = get_conn()
        try:
            with cn.cursor() as cur:
                units = active_units(cur)
        finally:
            cn.close()
    except Exception:
        log.warning("meta: อ่านหน่วยจากฐานข้อมูลไม่ได้ ใช้หน่วยเริ่มต้นแทน", exc_info=True)
        units = DEFAULT_UNITS
    return {"ok": True, "today": now_th().date().isoformat(), "units": units,
            "job_prefix": JOB_PREFIX, "job_digits": JOB_DIGITS}


@app.get("/api/branches")
def branches():
    # ส่งรายชื่อจาก All Shop ให้หน้าเว็บค้นหาได้ทันทีขณะพิมพ์
    cn = get_conn()
    try:
        with cn.cursor() as cur:
            cur.execute("SELECT brand, shop_id, code, name, alternate_name, tags FROM transport_notice_branches "
                        "WHERE is_active=1 AND hidden=0 ORDER BY brand, code")
            items = [{**r, "tags": split_tags(r["tags"])} for r in cur.fetchall()]
    finally:
        cn.close()
    return {"ok": True, "items": items}


# =============================================================================
#  หน้าตั้งค่า (ปุ่มรูปรถมุมซ้ายบน): เพิ่ม/แก้/ซ่อนสาขา + แท็ก, แก้หน่วยที่ส่ง
#  ไม่มี PIN — ทุกคนที่เปิดเว็บได้ใช้ได้ (ผู้ใช้เลือก 20 ต.ค. 2026)
# =============================================================================
def split_tags(text):
    return [t for t in (text or "").split(",") if t]


def norm_tags(tags):
    out, seen = [], set()
    for t in tags or []:
        t = " ".join(str(t).replace(",", " ").split())
        if t and t.casefold() not in seen:
            if len(t) > 30:
                raise HTTPException(422, f"แท็กยาวเกิน 30 ตัวอักษร: {t}")
            seen.add(t.casefold())
            out.append(t)
    joined = ",".join(out)
    if len(joined) > 255:
        raise HTTPException(422, "แท็กรวมกันยาวเกินไป")
    return joined


def active_units(cur):
    cur.execute("SELECT name FROM transport_notice_units WHERE is_active=1 ORDER BY sort_order, id")
    return [r["name"] for r in cur.fetchall()] or DEFAULT_UNITS


class BranchIn(BaseModel):
    model_config = ConfigDict(str_strip_whitespace=True)
    code: str = Field(..., min_length=1, max_length=20, pattern=r"^[0-9A-Za-z._-]+$")
    name: str = Field(..., min_length=1, max_length=255)
    alternate_name: str = Field("", max_length=255)
    tags: list[str] = []


class BranchUpdate(BaseModel):
    model_config = ConfigDict(str_strip_whitespace=True)
    name: Optional[str] = Field(None, min_length=1, max_length=255)
    alternate_name: Optional[str] = Field(None, max_length=255)
    tags: Optional[list[str]] = None
    hidden: Optional[bool] = None


class UnitIn(BaseModel):
    model_config = ConfigDict(str_strip_whitespace=True)
    name: str = Field(..., min_length=1, max_length=10)


class UnitUpdate(BaseModel):
    model_config = ConfigDict(str_strip_whitespace=True)
    name: Optional[str] = Field(None, min_length=1, max_length=10)
    hidden: Optional[bool] = None


class UnitMove(BaseModel):
    dir: int = Field(..., ge=-1, le=1)


BRANCH_SETTING_COLS = "id, brand, shop_id, code, name, alternate_name, tags, hidden, is_manual"


def _branch_out(row):
    return {**row, "tags": split_tags(row["tags"]), "hidden": bool(row["hidden"]), "is_manual": bool(row["is_manual"])}


@app.get("/api/settings/branches")
def settings_branches():
    cn = get_conn()
    try:
        with cn.cursor() as cur:
            cur.execute(f"SELECT {BRANCH_SETTING_COLS} FROM transport_notice_branches WHERE is_active=1 ORDER BY code")
            return {"ok": True, "items": [_branch_out(r) for r in cur.fetchall()]}
    finally:
        cn.close()


@app.post("/api/settings/branches")
def add_branch(body: BranchIn):
    tags = norm_tags(body.tags)
    cn = get_conn()
    try:
        with cn.cursor() as cur:
            cur.execute("SELECT id, is_active FROM transport_notice_branches WHERE brand=%s AND shop_id=%s",
                        (MANUAL_BRAND, body.code))
            old = cur.fetchone()
            if old and old["is_active"]:
                raise HTTPException(409, f"รหัส {body.code} มีอยู่แล้ว — ค้นหาแล้วกดแก้ไขแทน")
            if old:   # เคยมีแต่ถูกปิดจากการนำเข้า → เปิดใหม่เป็นสาขาที่เพิ่มเอง
                cur.execute("UPDATE transport_notice_branches SET code=%s, name=%s, alternate_name=%s, tags=%s, "
                            "is_active=1, hidden=0, is_manual=1 WHERE id=%s",
                            (body.code, body.name, body.alternate_name, tags, old["id"]))
                new_id = old["id"]
            else:
                cur.execute("INSERT INTO transport_notice_branches (brand, shop_id, code, name, alternate_name, tags, is_manual) "
                            "VALUES (%s,%s,%s,%s,%s,%s,1)",
                            (MANUAL_BRAND, body.code, body.code, body.name, body.alternate_name, tags))
                new_id = cur.lastrowid
            cn.commit()
            cur.execute(f"SELECT {BRANCH_SETTING_COLS} FROM transport_notice_branches WHERE id=%s", (new_id,))
            return {"ok": True, "item": _branch_out(cur.fetchone())}
    finally:
        cn.close()


@app.put("/api/settings/branches/{branch_id}")
def update_branch(branch_id: int, body: BranchUpdate):
    changes = body.model_dump(exclude_none=True)
    if "tags" in changes:
        changes["tags"] = norm_tags(changes["tags"])
    if "hidden" in changes:
        changes["hidden"] = int(changes["hidden"])
    if not changes:
        raise HTTPException(422, "ไม่มีข้อมูลที่จะแก้")
    cn = get_conn()
    try:
        with cn.cursor() as cur:
            sets = ", ".join(f"{k}=%s" for k in changes)
            cur.execute(f"UPDATE transport_notice_branches SET {sets} WHERE id=%s AND is_active=1",
                        [*changes.values(), branch_id])
            cn.commit()
            cur.execute(f"SELECT {BRANCH_SETTING_COLS} FROM transport_notice_branches WHERE id=%s AND is_active=1",
                        (branch_id,))
            row = cur.fetchone()
            if not row:
                raise HTTPException(404, "ไม่พบสาขานี้")
            return {"ok": True, "item": _branch_out(row)}
    finally:
        cn.close()


def _units_out(cur):
    cur.execute("SELECT id, name, sort_order, is_active FROM transport_notice_units ORDER BY sort_order, id")
    return [{"id": r["id"], "name": r["name"], "hidden": not r["is_active"]} for r in cur.fetchall()]


@app.get("/api/settings/units")
def settings_units():
    cn = get_conn()
    try:
        with cn.cursor() as cur:
            return {"ok": True, "items": _units_out(cur)}
    finally:
        cn.close()


@app.post("/api/settings/units")
def add_unit(body: UnitIn):
    cn = get_conn()
    try:
        with cn.cursor() as cur:
            cur.execute("SELECT id, is_active FROM transport_notice_units WHERE name=%s", (body.name,))
            old = cur.fetchone()
            if old and old["is_active"]:
                raise HTTPException(409, f"มีหน่วย “{body.name}” อยู่แล้ว")
            if old:
                cur.execute("UPDATE transport_notice_units SET is_active=1 WHERE id=%s", (old["id"],))
            else:
                cur.execute("SELECT COALESCE(MAX(sort_order), 0) + 1 AS n FROM transport_notice_units")
                cur.execute("INSERT INTO transport_notice_units (name, sort_order) VALUES (%s,%s)",
                            (body.name, cur.fetchone()["n"]))
            cn.commit()
            return {"ok": True, "items": _units_out(cur)}
    finally:
        cn.close()


@app.put("/api/settings/units/{unit_id}")
def update_unit(unit_id: int, body: UnitUpdate):
    cn = get_conn()
    try:
        with cn.cursor() as cur:
            cur.execute("SELECT id, is_active FROM transport_notice_units WHERE id=%s", (unit_id,))
            unit = cur.fetchone()
            if not unit:
                raise HTTPException(404, "ไม่พบหน่วยนี้")
            if body.name is not None:
                cur.execute("SELECT id FROM transport_notice_units WHERE name=%s AND id<>%s", (body.name, unit_id))
                if cur.fetchone():
                    raise HTTPException(409, f"มีหน่วย “{body.name}” อยู่แล้ว")
                cur.execute("UPDATE transport_notice_units SET name=%s WHERE id=%s", (body.name, unit_id))
            if body.hidden is not None:
                if body.hidden and unit["is_active"]:
                    cur.execute("SELECT COUNT(*) AS n FROM transport_notice_units WHERE is_active=1")
                    if cur.fetchone()["n"] <= 1:
                        raise HTTPException(409, "ต้องเหลือหน่วยที่ใช้งานอย่างน้อย 1 หน่วย")
                cur.execute("UPDATE transport_notice_units SET is_active=%s WHERE id=%s", (int(not body.hidden), unit_id))
            cn.commit()
            return {"ok": True, "items": _units_out(cur)}
    except Exception:
        cn.rollback()
        raise
    finally:
        cn.close()


@app.post("/api/settings/units/{unit_id}/move")
def move_unit(unit_id: int, body: UnitMove):
    cn = get_conn()
    try:
        with cn.cursor() as cur:
            cur.execute("SELECT id FROM transport_notice_units ORDER BY sort_order, id")
            ids = [r["id"] for r in cur.fetchall()]
            if unit_id not in ids:
                raise HTTPException(404, "ไม่พบหน่วยนี้")
            i = ids.index(unit_id)
            j = i + body.dir
            if 0 <= j < len(ids):
                ids[i], ids[j] = ids[j], ids[i]
                cur.executemany("UPDATE transport_notice_units SET sort_order=%s WHERE id=%s",
                                [(n, uid) for n, uid in enumerate(ids, 1)])
                cn.commit()
            return {"ok": True, "items": _units_out(cur)}
    finally:
        cn.close()


@app.get("/api/jobs")
def list_jobs(date: Optional[datetime.date] = Query(None, description="วันที่บันทึก (ค่าเริ่มต้น = วันนี้)")):
    work_date = date or now_th().date()
    cn = get_conn()
    try:
        with cn.cursor() as cur:
            cur.execute(JOB_SELECT + " WHERE work_date=%s AND deleted_at IS NULL ORDER BY seq",
                        (work_date,))
            items = [_out(r) for r in cur.fetchall()]
    finally:
        cn.close()
    return {"ok": True, "date": work_date.isoformat(), "items": items}


@app.post("/api/jobs")
def create_job(body: JobIn):
    now = now_th()
    work_date = now.date()
    data = body.model_dump()
    cn = get_conn()
    try:
        with cn.cursor() as cur:
            units = active_units(cur)
            if data["unit"] not in units:
                raise HTTPException(422, "หน่วยต้องเป็น " + "/".join(units))
            # จองเลขงานของวันแบบ atomic (แถวตัวนับถูก lock จน commit → ไม่ได้เลขซ้ำแม้กดพร้อมกัน)
            cur.execute(
                "INSERT INTO transport_notice_daily_counter (work_date, last_seq) VALUES (%s, LAST_INSERT_ID(1)) "
                "ON DUPLICATE KEY UPDATE last_seq = LAST_INSERT_ID(last_seq + 1)",
                (work_date,),
            )
            cur.execute("SELECT LAST_INSERT_ID() AS seq")
            seq = cur.fetchone()["seq"]

            cols = ["work_date", "seq", "job_no", *WRITE_FIELDS, "created_at", "updated_at"]
            vals = [work_date, seq, job_no(seq), *[data[f] for f in WRITE_FIELDS], now, now]
            cur.execute(
                f"INSERT INTO transport_notice_jobs ({', '.join(cols)}) VALUES ({', '.join(['%s'] * len(cols))})",
                vals,
            )
            new_id = cur.lastrowid
        cn.commit()
        with cn.cursor() as cur:
            return {"ok": True, "item": _get_job(cur, new_id)}
    except Exception:
        cn.rollback()
        raise
    finally:
        cn.close()


# รายการที่บันทึกแล้วแก้ไขไม่ได้ (ตามที่ผู้ใช้กำหนด 20 ต.ค. 2026) — จะแก้ให้คัดลอกเป็นรายการใหม่แล้วลบรายการเดิม
# คง route ไว้เพื่อตอบข้อความชัดเจนให้หน้าเว็บเวอร์ชันเก่าที่ยังค้างในเบราว์เซอร์
@app.put("/api/jobs/{job_id}")
def update_job(job_id: int):
    raise HTTPException(403, "รายการที่บันทึกแล้วแก้ไขไม่ได้ — กด “คัดลอก” แก้ข้อมูลแล้วบันทึก จากนั้นลบรายการเดิม")


@app.delete("/api/jobs/{job_id}")
def delete_job(job_id: int):
    cn = get_conn()
    try:
        with cn.cursor() as cur:
            _get_job(cur, job_id)
            cur.execute("UPDATE transport_notice_jobs SET deleted_at=%s WHERE id=%s", (now_th(), job_id))
        cn.commit()
    finally:
        cn.close()
    return {"ok": True}


@app.get("/api/export")
def export_jobs(date: Optional[datetime.date] = None, q: str = ""):
    """Export real XLSX without loading third-party JavaScript in the browser."""
    work_date = date or now_th().date()
    rows = list_jobs(work_date)["items"]
    search_fields = ["job_no", "to_name", "to_dept", "to_branch_code", "to_branch_name",
                     "from_name", "from_dept", "from_branch_code", "from_branch_name",
                     "item_detail", "unit", "remark"]
    tokens = q.lower().split()
    rows = [r for r in rows if all(t in " ".join(str(r.get(f, "")) for f in search_fields).lower()
                                   for t in tokens)]
    filename = quote(f"ใบแจ้งงานขนส่ง_{work_date}.xlsx")
    return Response(build_xlsx(rows), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"})


def build_xlsx(rows):
    """ไฟล์ Excel 14 คอลัมน์ตามฟอร์ม SST — ใช้ทั้งปุ่ม Excel บนเว็บและไฟล์แนบอีเมลรายวัน"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    columns = [("ลำดับ", "seq", 8), ("รหัสงาน Job No.", "job_no", 18),
               ("วันที่ฝาก (ว/ด/ป)", "deposit_date", 20), ("วันที่ส่ง (ว/ด/ป)", "delivery_date", 20),
               ("ถึง ผู้รับ", "to_name", 24), ("แผนก ผู้รับ", "to_dept", 24),
               ("สาขา ผู้รับ", "to_branch_code", 22), ("จาก ผู้ส่ง", "from_name", 24),
               ("แผนก ผู้ส่ง", "from_dept", 24), ("สาขา ผู้ส่ง", "from_branch_code", 22),
               ("รายละเอียด สิ่งที่ส่ง", "item_detail", 40), ("จำนวน", "qty", 12),
               ("หน่วย", "unit", 12), ("หมายเหตุ", "remark", 35)]
    wb = Workbook()
    ws = wb.active
    ws.title = "ใบแจ้งงานขนส่ง"
    ws.append([c[0] for c in columns])
    for index, (_, _, width) in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
        cell = ws.cell(1, index)
        cell.font = Font(name="Tahoma", bold=True, color="FFFFFF", size=12)
        cell.fill = PatternFill("solid", fgColor="4F6B33" if 8 <= index <= 10 else "1B2E4F")
    for row_number, row in enumerate(rows, 2):
        for index, (_, key, _) in enumerate(columns, 1):
            value = row.get(key, "")
            if key in ("deposit_date", "delivery_date"):
                value = datetime.date.fromisoformat(str(value))
            cell = ws.cell(row_number, index, value)
            # User-entered text beginning with '=' remains literal text.
            if isinstance(value, str):
                cell.data_type = "s"
            cell.font = Font(name="Tahoma", size=12)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(value, datetime.date):
                cell.number_format = "d/m/yyyy"
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:1"
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


# Serve the actual root index without exposing project files or .env.
@app.get("/", include_in_schema=False)
@app.get("/index.html", include_in_schema=False)
def web_index():
    return FileResponse(INDEX_FILE, media_type="text/html", headers={"Cache-Control": "no-cache"})


@app.get("/docs/index.html", include_in_schema=False)
def legacy_index():
    return RedirectResponse("/", status_code=307)


if os.path.isdir(DOCS_DIR):
    app.mount("/docs", StaticFiles(directory=DOCS_DIR), name="assets")
    # Existing tabs may still request the old asset URLs until they reload.
    app.mount("/css", StaticFiles(directory=os.path.join(DOCS_DIR, "css")), name="legacy-css")
    app.mount("/js", StaticFiles(directory=os.path.join(DOCS_DIR, "js")), name="legacy-js")
